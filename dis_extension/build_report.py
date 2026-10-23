"""Build the evidence-linked report after all held-out measurements and audits."""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import ast
import json
import sys
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
sys.path.insert(0,str(ROOT))
from build_manuscript import Manuscript, wrap_code
from native_equations import prepare_equations
from transfer_reporting import abstract as transfer_abstract, append_sections as append_transfer_sections

FORMULAS={
 'problem':r'V(c)=\max_{x\geq0,\ Ax\leq c}\sum_iw_i\log(x_i+\xi_i),\qquad W=\sum_iw_i.',
 'certificate':r'0\leq V(c)-U(x)\leq B(c,\lambda)-U(x),\qquad \frac{B(c,\lambda)-U(x)}{W}\leq\epsilon.',
 'incumbent':r'\overline B_k=\min_{j\leq k}B(c,\lambda_j),\qquad \underline U_k=\max_{j\leq k}U(x_j).',
 'ratio':r'\lambda^{\rm ratio}_e=\max\left\{10^{-12},\lambda_{0,e}\frac{c_{0,e}}{c_e}\right\}.',
 'hessian':r'H_{SS}=A_S\operatorname{diag}\left(\frac{w_i}{q_i^2}\mathbf{1}_{\{q_i\xi_i<w_i\}}\right)A_S^\top,\qquad q=A^\top\lambda.',
 'target':r'y(z)=\log\frac{T_{\rm refine}}{T_{\rm reuse}},\qquad \pi(z)=\mathbf{1}_{\{f(z)<\tau\}},\quad\tau\in\{0,-0.05\}.',
}
TITLE='Learning to Refine Certified Network Allocations'
LABELS={'ray_warm':'Ray warm start','ray_reuse':'Original reuse','always_refine':'Always refine',
        'heuristic':'Capacity heuristic','train_best':'Training best fixed','learned':'Learned selector'}


def wrap_excerpt(source):
    lines=[]
    for line in wrap_code(source,width=88).splitlines():
        if len(line)>88 and ' if ' in line and ' else ' in line:
            before,rest=line.split(' if ',1)
            condition,after=rest.split(' else ',1)
            continuation=' '*(len(line)-len(line.lstrip())+4)
            lines.extend([before,continuation+'if '+condition,continuation+'else '+after])
        else:
            lines.append(line)
    result='\n'.join(lines)+'\n'
    assert ast.dump(ast.parse(source),include_attributes=False)==ast.dump(ast.parse(result),include_attributes=False)
    assert max(map(len,lines))<=95
    return result


class Report(Manuscript):
    def __init__(self,metadata):
        super().__init__(metadata)
        props=self.doc.core_properties
        props.title=TITLE;props.subject='Cost-sensitive learning, grouped validation and certified numerical updates'
        props.keywords='scientific computing; algorithm selection; grouped validation; network utility maximisation'
        props.created=datetime.now(timezone.utc);props.modified=props.created
        header=self.doc.sections[0].header.paragraphs[0]
        header.clear()
        run=header.add_run(TITLE);run.font.name='Times New Roman';run.font.size=Pt(8)
        run.italic=True;run.font.color.rgb=RGBColor.from_string('606060')

    def eq(self,key):
        self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
        nodes=self.meta[key]
        for i,node in enumerate(nodes):
            p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5 if i==0 else 0)
            p.paragraph_format.space_after=Pt(8 if i==len(nodes)-1 else 2)
            p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=i<len(nodes)-1
            p._p.append(deepcopy(node))
        self.transcript.append('\n$$'+FORMULAS[key]+'$$\n')

    def figure(self,name,caption,width=6.55):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(HERE/'figures'/(name+'.png')),width=Inches(width))
        cp=self.doc.add_paragraph(style='Caption');label,body=caption.split('. ',1)
        cp.add_run(label+'. ').bold=True;cp.add_run(body)
        self.transcript.append('\n![Scientific figure](../dis_extension/figures/'+name+'.png)\n'+caption+'\n')


def main():
    folder=HERE/'results'
    has_transfer=(HERE/'transfer_results/summary.json').exists()
    summary=json.loads((folder/'summary.json').read_text())
    audit=json.loads((folder/'audit.json').read_text());assert audit['passed']
    training=json.loads((folder/'training_environment.json').read_text())
    core=summary['paired']['core']['learned'];shift=summary['paired']['shift']['learned']
    coverage=summary['methods']; rcore=core['vs_reuse']['median_ratio'];rshift=shift['vs_reuse']['median_ratio']
    m=Report(prepare_equations(FORMULAS,ROOT/'tmp/dis_equations'))
    m.doc.add_paragraph('Learning to Refine\nCertified Network Allocations','Title')
    m.doc.add_paragraph('Cost sensitive selection with topology held out validation','Subtitle')
    m.transcript.append('# '+TITLE+'\n')
    p=m.p('Kai Chi\n'+('23 October 2026' if has_transfer else '22 October 2026'),size=10.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(12)
    m.sub('Abstract')
    if has_transfer:
        m.p(transfer_abstract())
    else:
        m.p(f"Repeated optimisation creates a computational decision: when is a lightweight correction worth its cost? This study extends a certified network utility maximisation pipeline with a capacity-ratio price predictor, two bounded coordinate Newton corrections and a learned choice between refinement and the original reuse action. The predictor selects computation; an actual feasible-primal/dual certificate determines acceptance. Historical action costs are collected on 576 capacity cases from six public topology structures. Model and threshold selection use nested leave-one-topology-out validation, with no updated reference solutions or test timings among online inputs. Fresh paired replays compare six pipelines on those cases and 192 additional cases with shifted capacity-change distributions, for 13,824 timing records. The learned selector's median paired runtime ratio to original reuse is {rcore:.3f} on core cases and {rshift:.3f} on shifted cases. Corresponding ratios to the strong ray warm start are {core['vs_ray_warm']['median_ratio']:.3f} and {shift['vs_ray_warm']['median_ratio']:.3f}. There are {summary['failed_runs']} replay runs that do not meet the target; they are retained and not accepted as certified outputs. A simple capacity rule matches core aggregate cost and performs better under shifted shocks. Learning does not consistently add value beyond the numerical update. All online feature, prediction, correction and fallback costs are charged. Six convenience topology families and synthetic workloads limit external inference.")
    m.sub('Research question')
    m.p('Can historical numerical observations guide a low-cost update decision across unseen network families, while every accepted allocation remains subject to the same explicit utility-error target? The contribution is an auditable applied study of this question; algorithm selection, duality and Newton updates are established tools.')

    m.heading('1 From a certificate limitation to a data question')
    m.p('The parent study considers hard-capacity allocation after a change in link capacities. It derives specialised shifted-log error bounds and exhibits a two-link case in which a feasible repaired allocation is optimal while a single price-ray certificate is arbitrarily loose. Its large benchmark, however, gives a different empirical diagnosis: 216 of 576 core candidates satisfy the primary tolerance and all are certified. For the other 360 candidates, the saved reference lower bound on utility loss exceeds the tolerance. A stronger upper bound alone cannot make those particular candidates meet the target.')
    m.p('This extension consequently improves candidate allocations and prices together, then asks whether a data-fitted policy can choose when that extra computation is worthwhile. The two-link example motivates escaping a fixed price direction. It does not, by itself, predict performance on coupled networks. The core data were already explored before this extension; the new action and evaluation protocol was fixed before the reported action collection.')
    m.p('The starting NUM literature includes accelerated optimisation [1], first-order multiplier adaptation after system changes [2], and scalable proximal message passing with capacity-degradation warm starts [3]. Per-instance algorithm selection is a well-developed research problem [4]. This report does not claim that price prediction or selecting an algorithm from historical observations is new.')

    m.heading('2 Model and acceptance invariant')
    m.p('Let A be a nonnegative link-route matrix with nonempty routes, c positive capacities, w positive utility weights and xi nonnegative shifts. Routes and utility parameters stay fixed within each update. Natural logarithms are used.')
    m.eq('problem')
    m.p('The dual upper bound B is obtained by maximising the separable Lagrangian over nonnegative rates. For any feasible allocation and valid nonnegative prices with positive route prices, weak duality gives the following acceptance rule, with epsilon equal to 0.001.')
    m.eq('certificate')
    m.p('The guarantee concerns weighted aggregate utility. It does not impose a minimum individual throughput. Every proposed allocation is repaired against the actual updated capacities. Predictions never declare a candidate acceptable: they select an action, which must then evaluate the certificate. If an inexpensive update fails the target, the common L-BFGS-B solver is invoked. A failed fallback is retained as a failed run.')
    m.sub('Retaining valid incumbents')
    m.p('Candidates and prices from different update steps can be combined by retaining the best feasible utility and the smallest valid dual upper bound. Each feasible candidate lies below the optimum and each dual value lies above it, so their difference remains a valid bound.')
    m.eq('incumbent')
    m.p('Adding valid incumbents cannot increase the retained gap. This is a standard duality consequence. It does not establish a runtime improvement, a new global convergence theorem or numerical validity under arbitrary rounding. The implementation uses Float64 feasibility guards, rather than interval arithmetic.')

    m.heading('3 Lightweight candidate and price correction')
    m.p('Both actions first run the original price-ray query and compare the repaired nominal candidate with the repaired ray-stationary candidate. If that certificate meets the target, both actions return immediately. Otherwise the original action calls the common solver, whereas the refinement action first proposes componentwise capacity-ratio prices.')
    m.eq('ratio')
    m.p('For independent links with unshifted logarithmic utility and an exact nominal solution, the optimal prices are w divided by capacity. The predictor therefore gives exact updated prices in that special case, including the parent study\'s two-link counterexample. On coupled networks it is a heuristic; it can worsen an individual proposal, which is why the preceding incumbents are retained.')
    m.p('If the target is still unmet, the method attempts at most two corrections on at most eight price coordinates. Eligible coordinates are ranked by absolute dual gradient divided by the square root of the regularised Hessian diagonal. A price at the positive floor is eligible only when the gradient points towards increasing that price. Within a fixed positive-rate active set, the selected Hessian is')
    m.eq('hessian')
    m.p('The reduced Newton system includes a ridge of 1e-8 times the maximum Hessian diagonal, with a positive lower guard. Prices are projected to a floor of 1e-12. An Armijo test tries at most eight step lengths. Accepted trial rates are repaired and the primal and dual incumbents are updated independently. Activity can change, so the local Hessian formula is used as an update mechanism with safeguarding, rather than a global approximation theorem. All proposal, linear-algebra and line-search costs are timed.')

    m.heading('4 Learning computational value')
    m.p('The selector is fitted on historical cases rejected by the initial certificate. Its target is the logarithm of the ratio between complete refinement-action and original-action runtimes; each runtime is a median of three observations. The fitted function and a threshold decide whether to refine.')
    m.eq('target')
    m.table(['Descriptor group','Information available before solving the new problem'],[
        ('Problem size','Flow count, link count and mean route length'),
        ('Nominal allocation','Weighted mean and spread of utility-shift fractions'),
        ('Capacity change','Minimum, mean, spread, maximum and fraction of changed capacities'),
        ('Nominal price weighting','Price-weighted capacity ratios and active-price fraction'),
        ('Initial certificate','Log gap, ray scale and candidate utility improvement'),
    ],[1.8,5.05],font=9.3)
    m.p('There are 16 numerical features. Topology names, case identifiers, shock labels, updated optimal values, updated reference prices, updated reference allocations and current-case action runtimes are excluded from online inputs. The target may use historical training runtimes; it is never supplied with the held-out case\'s measured action costs during inference. Initial certification is checked before feature extraction, so already certified cases do not incur model inference.')
    m.p('The candidate set contains ridge regression with penalties 1 and 10, and regression trees with depths 1, 2 and 3 and a minimum of 20 training examples per leaf. The decision threshold is selected from 0 and -0.05 in log-ratio units. A fitted ridge model or tree is exported as JSON and evaluated directly with numerical arrays or tree traversal. The deployed prediction path does not invoke an estimator object or access reference data.')
    m.sub('Nested topology held out validation')
    m.p('Each outer fold excludes one complete topology family, including every seed, size, capacity law, utility shift and shock associated with it. The remaining five topologies supply training data. An inner leave-one-topology-out loop chooses the model specification and decision threshold. Its score gives equal weight to each validation topology and includes measured feature-extraction overhead. Prediction cost is measured in the final replay rather than estimated separately in this selection score. Standardisation is fitted within each training split. These controls address the grouped-data and preprocessing leakage issues discussed in [5,6].')
    m.p('Only the 360 initially rejected core cases are used for fitting, giving 300 fitting cases in each outer fold before inner splitting. All 96 core cases of the excluded topology are evaluated, including its automatically accepted cases. A separate all-core-data model is supplied for future use; first-phase held-out results use the outer-fold models. The training-best fixed-action baseline is also chosen from training topologies alone.')

    m.heading('5 First phase data and measurement protocol')
    m.p('The six topology structures are Abilene, Geant2012, Cernet, Colt, UsCarrier and Cogentco, pinned to the public Internet Topology Zoo sources and licence supplied with the parent study [7]. They are historical structures, rather than observations of current network operation. Routes, weights, capacities and capacity changes are generated. Workloads contain 1,000 or 10,000 flows; this extension does not claim a new 100,000-flow evaluation.')
    m.table(['Stage','Cases and observations','Role'],[
        ('Action collection','576 cases; 3,456 timings','Historical costs for two actions'),
        ('Outer validation','6 complete topology folds','Transfer across graph families'),
        ('Shifted shocks','192 cases on 48 nominal anchors','New capacity-change distributions'),
        ('Paired replay','768 cases; 13,824 timings','End-to-end evaluation of six policies'),
    ],[1.35,2.65,2.85],font=9.2)
    m.p('The shifted cases use nominal anchors with seed 101, both flow counts, both capacity laws and both utility shifts. Their four shock types are 5% reduction on a random 5% of links; 50% reduction on a random 5% of links; independent reductions up to 60%; and independent capacity factors between 0.8 and 1.2. The last setting includes both losses and recovery. No shifted-case result is used for fitting or selection. Because nominal anchors are reused, these are new shocks rather than 192 new independent networks.')
    m.table(['Pipeline','Online computation and decision'],[
        ('Ray warm start','Direct common solve from scalar-adjusted prices; no candidate precheck'),
        ('Original reuse','Run the common solver from the original query output'),
        ('Always refine','Apply the bounded correction before any necessary fallback'),
        ('Capacity heuristic','Refine when the standard deviation of c/c0 exceeds 1e-4'),
        ('Training best fixed','Choose one fixed action using training-topology costs'),
        ('Learned selector','Use the fitted outer-fold model to choose the action'),
    ],[1.65,5.2],font=9.2)
    m.p('Methods are freshly replayed three times per case with rotating execution order. Data loading and offline training are outside online timers. Feature extraction, prediction, candidate repair, certificates, refinement, line search and fallback are inside them. A common cache preparation cost is recorded separately and amortised over four updates in a secondary field. The primary runtime comparison uses complete online costs. Numerical-library thread counts are checked to be one. An initial collection with inconsistent NumPy and SciPy thread settings was interrupted; its observations are retained separately and excluded from training, selection, evaluation and figures.')
    m.p(f"Offline nested selection and fitting took {training['training_seconds']:.2f} seconds in the recorded environment. Timing repetitions measure execution variation, not independent network draws. Per-case medians are the basic runtime observations. Topology-block bootstrap intervals use 10,000 resamples of the six families; they describe this collection and are not population-coverage claims after model selection.")

    m.heading('6 First phase computational performance')
    if summary['failed_runs']:
        m.p('The primary timing summaries retain runs that reach the computational cap without certification. Their cost is not equivalent to successfully obtaining an allocation at the target accuracy. The success counts below and a descriptive common-success comparison are therefore reported alongside runtime ratios.')
        m.table(['Pipeline','Core solved cases','Shift solved cases'],[
            (label,str(coverage['core'][method]['solved'])+' / 576',
             str(coverage['shift'][method]['solved'])+' / 192')
            for method,label in LABELS.items()],[2.55,2.15,2.15],font=9.2)
    m.figure('01_held_out_runtime','Figure 1. Paired runtime medians relative to original reuse, with descriptive topology-block bootstrap intervals. Online feature, inference, correction and fallback costs are included. Values below one favour the displayed pipeline.')
    rows=[]
    for method in LABELS:
        a=summary['paired']['core'][method]['vs_reuse'];b=summary['paired']['shift'][method]['vs_reuse']
        rows.append((LABELS[method],f"{a['median_ratio']:.3f}",f"{b['median_ratio']:.3f}",f"{a['total_time_ratio']:.3f} / {b['total_time_ratio']:.3f}"))
    m.table(['Pipeline','Core median ratio','Shift median ratio','Total time ratio core / shift'],rows,[1.9,1.45,1.45,2.05],font=9.0)
    m.p(f"The learned selector's core median ratio is {rcore:.3f}, with a topology-block interval [{core['vs_reuse']['topology_bootstrap_interval'][0]:.3f}, {core['vs_reuse']['topology_bootstrap_interval'][1]:.3f}]. Its shifted-case ratio is {rshift:.3f}, with interval [{shift['vs_reuse']['topology_bootstrap_interval'][0]:.3f}, {shift['vs_reuse']['topology_bootstrap_interval'][1]:.3f}]. Median paired ratios and total-time ratios answer different questions: a favourable median can coexist with an expensive tail or a small aggregate gain.")
    if summary['failed_runs']:
        a=summary['common_success_paired']['core']['learned']
        b=summary['common_success_paired']['shift']['learned']
        m.p(f"Restricting to cases certified in all repetitions by all six pipelines leaves {a['cases']} core and {b['cases']} shifted cases. The corresponding learned-to-original median ratios are {a['median_ratio']:.3f} and {b['median_ratio']:.3f}. This conditional description does not replace the full-workload failure counts or establish performance on unresolved cases.")
    m.p(f"Against the training-selected fixed action, the learned selector's median ratios are {core['vs_train_best']['median_ratio']:.3f} on core cases and {shift['vs_train_best']['median_ratio']:.3f} on shifted cases. Against the strong ray warm start, they are {core['vs_ray_warm']['median_ratio']:.3f} and {shift['vs_ray_warm']['median_ratio']:.3f}. These comparisons are necessary to distinguish the value of the numerical correction from the additional value of learning when to use it.")
    heuristic_core=summary['paired']['core']['heuristic']['vs_reuse']
    heuristic_shift=summary['paired']['shift']['heuristic']['vs_reuse']
    m.p(f"The simple capacity heuristic is a particularly informative comparator. Its core median ratio is {heuristic_core['median_ratio']:.3f} and its total-time ratio is {heuristic_core['total_time_ratio']:.3f}, compared with {rcore:.3f} and {core['vs_reuse']['total_time_ratio']:.3f} for learning. Under shifted shocks the heuristic's ratios are {heuristic_shift['median_ratio']:.3f} and {heuristic_shift['total_time_ratio']:.3f}, compared with {rshift:.3f} and {shift['vs_reuse']['total_time_ratio']:.3f}. Learning therefore does not demonstrate an overall computational advantage over this simple rule. The small core differences should not be interpreted as a stable ordering across hardware or repeated execution environments.")
    rejected_core=summary['rejected_paired']['core']['learned']
    rejected_shift=summary['rejected_paired']['shift']['learned']
    m.p(f"A descriptive secondary comparison restricts evaluation to cases rejected by the common initial certificate, where a policy decision is required. This comprises {rejected_core['vs_reuse']['cases']} core cases and {rejected_shift['vs_reuse']['cases']} shifted cases. The learned-to-original median ratios are {rejected_core['vs_reuse']['median_ratio']:.3f} and {rejected_shift['vs_reuse']['median_ratio']:.3f}; the learned-to-training-fixed ratios are {rejected_core['vs_train_best']['median_ratio']:.3f} and {rejected_shift['vs_train_best']['median_ratio']:.3f}. The subset is determined by the common certificate, independently of observed policy speed. No model is refitted for this analysis.")
    m.p(f"Within that rejected core subset, the learned-to-ray-warm median ratio is {rejected_core['vs_ray_warm']['median_ratio']:.3f}. The favourable all-core comparison with ray warm start thus also reflects the shared initial certificate, which avoids full solves on already accepted cases; it should not be attributed entirely to the learned decision.")
    m.p(f"For cases requiring a decision, the median measured feature-plus-prediction overhead is {summary['inference_overhead_rejected_us']['core']:.1f} microseconds in core evaluation and {summary['inference_overhead_rejected_us']['shift']:.1f} microseconds under shifted shocks. Offline fitting is a separate cost and is not represented as free. Changes in hardware, shared-host load or workload mix may alter the value of a gate fitted on historical runtime ratios.")

    m.heading('7 Candidate improvement and distribution shift')
    m.figure('02_certificate_coverage','Figure 2. Cases certified by the initial query and additional cases certified after the lightweight update, before full L-BFGS-B. The added category includes numerical correction work and is not described as iteration-free reuse.')
    m.p(f"In the core workload, the initial certificate accepts {coverage['core']['ray_reuse']['precheck_certified']} cases. Always applying refinement certifies another {coverage['core']['always_refine']['light_certified']} before full fallback; the learned selector certifies another {coverage['core']['learned']['light_certified']}. In the shifted workload the corresponding initial and additional counts are {coverage['shift']['ray_reuse']['precheck_certified']}, {coverage['shift']['always_refine']['light_certified']} and {coverage['shift']['learned']['light_certified']}. A higher early-certification count is an intermediate outcome; total measured computation determines its practical value.")
    m.figure('03_topology_transfer','Figure 3. Learned-to-original paired runtime ratio within each held-out topology family. Core and shifted-shock workloads are shown separately; each point is descriptive for that family, not an independent population estimate.')
    m.p('The outer models may select different regression families and thresholds because each is fitted without its test topology. Training-set selection scores are supplied for every candidate. The comparison should be interpreted together with family-level results: an average over these six graphs can hide a transfer failure on one graph, and the same nominal anchor can produce several correlated cases.')
    m.figure('04_shock_transfer','Figure 4. Learned-to-original paired runtime medians by capacity-change type. These settings were fixed before the reported evaluation; all cases, including unfavourable ones, remain in the tables and raw observations.')
    m.p('The first-phase shifted workload changes severity and introduces recovery, while retaining the fixed routes and utility family. It does not test previously unseen routing policies, production traces, new graph families beyond the six outer folds, or a sequentially evolving cache. Those are separate distribution shifts. A model can continue to return certified solutions through fallback while losing its runtime advantage; validity and computational transfer must be assessed separately.')

    m.heading('8 Numerical audit and research interpretation')
    m.p(f"The fresh replay records {summary['failed_runs']} failed runs. The maximum observed relative overload is {summary['maximum_overload']:.2g}, and the maximum gap of a run marked converged is {summary['maximum_accepted_gap']:.6g}. Saved states for all 1,152 historical action outputs were recomputed independently of timing summaries. All 768 held-out policy choices were reconstructed from permitted features and their fold-specific models. Source hashes, numerical-library thread counts and topology/anchor separation passed the recorded audit.")
    m.p('All 18 failed runs come from one shifted case: UsCarrier, 10,000 flows, independent nominal capacities, unshifted log utility and heterogeneous reductions up to 60 percent. Each of the six pipelines reaches 3,000 accepted L-BFGS-B iterations in all three repetitions. The resulting normalised gaps range from 0.00249 to 0.00310, above the 0.001 target; no such output is accepted. All 576 core cases and the other 191 shifted cases are certified by every pipeline in every repetition. The shifted tight reference, which permits additional safeguarded Newton refinement, meets its separate target.')
    m.p('The strongest defensible conclusion concerns the measured trade-off between a bounded numerical correction and a cost-sensitive selection rule. The weak-duality invariant explains why a poor model choice changes computation rather than the acceptance target. The six-family transfer experiment evaluates whether observed runtime relationships persist under a deliberately separated topology and shock setting. Neither component establishes that a fitted model will accelerate arbitrary future scientific workloads.')
    if core['vs_train_best']['median_ratio']>=1 or shift['vs_train_best']['median_ratio']>=1:
        m.p(f"The learned selector's median ratio to the training-selected fixed action changes from {core['vs_train_best']['median_ratio']:.3f} on core cases to {shift['vs_train_best']['median_ratio']:.3f} under shifted shocks. The simple capacity rule matches core aggregate performance and performs better on the shifted workload. The results support measuring the numerical correction's value, while limiting a claim that learning is needed to obtain it. Feature and prediction costs, limited training groups and changing runtime relationships are plausible constraints; this experiment does not isolate their individual causal effects.")
    else:
        m.p('The learned selector improves median runtime relative to the training-selected fixed action in both reported workloads. This is evidence for adaptive selection within the specified collection. It is not a guarantee across other graph families, machines or capacity distributions, and it must be considered alongside total-time and family-level results.')
    m.p('Further work should obtain independent topology and workload collections, repeated execution environments and genuinely sequential updates with explicit cache-refresh costs. More expressive policies should be justified by measured net value over simple selectors. A matched-model modern conic or proximal-message-passing comparison would strengthen algorithmic positioning. The present six pipelines compare choices within the shared L-BFGS-B framework; they do not establish state-of-the-art solver superiority.')
    m.sub('Reproducibility')
    m.p('The companion archive contains the original source arrays, new capacities and reference arrays, all action and replay observations, explicit split manifests, fitted portable models, numerical checks and figure scripts. Model-selection and execution stages reject accidental overwrite of measured results. The complete pipeline can be rerun with pinned dependencies. Original public graph files and their CC BY 4.0 attribution are retained. No publication priority, external replication or journal acceptance is asserted.')

    if has_transfer:
        append_transfer_sections(m)
    m.heading('References',newpage=not has_transfer)
    refs=[
        ('Y. Tian et al. Centralized Network Utility Maximization with Accelerated Gradient Method. IEEE ICNP, 2022; author preprint arXiv:2408.08034.','https://arxiv.org/abs/2408.08034'),
        ('I.-H. Hou. Network Optimization in Dynamic Systems: Fast Adaptation via Zero-Shot Lagrangian Update. arXiv:2412.07865; accepted for IEEE INFOCOM 2025.','https://arxiv.org/abs/2412.07865'),
        ('A. Sreekumar, A. Degleris and R. Rajagopal. Large-Scale Network Utility Maximization via GPU-Accelerated Proximal Message Passing. IEEE Control Systems Letters 9, 2807-2812, 2025; author version arXiv:2509.10722v2, 2026.','https://arxiv.org/abs/2509.10722v2'),
        ('B. Bischl et al. ASlib: A Benchmark Library for Algorithm Selection. Artificial Intelligence 237, 41-58, 2016; author version arXiv:1506.02465v3.','https://arxiv.org/abs/1506.02465v3'),
        ('Scikit-learn developers. Cross-validation: evaluating estimator performance. Version 1.8.0 documentation; accessed 22 October 2026.','https://scikit-learn.org/1.8/modules/cross_validation.html'),
        ('Scikit-learn developers. Common pitfalls and recommended practices. Version 1.8.0 documentation; accessed 22 October 2026.','https://scikit-learn.org/1.8/common_pitfalls.html'),
        ('S. Knight et al. The Internet Topology Zoo. IEEE Journal on Selected Areas in Communications 29(9), 1765-1775, 2011. Graph source commit 51118849b99534f01010d885d66321433a7eef9e.','https://doi.org/10.1109/JSAC.2011.111002'),
    ]
    for i,(text,url) in enumerate(refs,1):m.reference(i,text,url)
    m.p('The parent manuscript, Certified Reuse for Network Utility Maximisation, supplies the shifted-log derivations, original benchmark and sharp ray-restriction counterexample. Its raw observations are retained in the companion archive. Methodological references above establish related work, not an assertion that their published timing results are comparable across machines.',size=10)
    tree=ast.parse((HERE/'certified_updates.py').read_text())
    nodes={n.name:n for n in tree.body if isinstance(n,ast.FunctionDef)}
    extracts=[('Features and portable policy inference','\n\n'.join(ast.unparse(nodes[k]) for k in ['features','predict'])),
              ('Certified action and fallback',ast.unparse(nodes['run_action']))]
    appendix=[]
    for letter,(title,source) in zip('AB',extracts):
        wrapped=wrap_excerpt(source);lines=wrapped.splitlines()
        for part,start in enumerate(range(0,len(lines),54)):
            label='Appendix '+letter+' '+title+(' continued' if part else '')
            m.heading(label,newpage=True)
            if not part:m.p('Extracted from certified_updates.py with AST-preserving presentation wrapping. Numerical correction, training, evaluation and audit code are supplied in full in the archive.',size=9.5)
            m.code('\n'.join(lines[start:start+54]));appendix.append({'heading':label,'lines':len(lines[start:start+54])})
    target=ROOT/'deliverables/DIS_Study.docx'
    target.parent.mkdir(parents=True,exist_ok=True)
    m.doc.save(target)
    (ROOT/'deliverables/DIS_Study.md').write_text('\n'.join(m.transcript))
    number_sources=['dis_extension/results/summary.json']
    if has_transfer:number_sources.append('dis_extension/transfer_results/summary.json')
    (folder/'report_build.json').write_text(json.dumps({'docx':str(target),'equation_blocks':len(FORMULAS),'appendix':appendix,'numbers_sources':number_sources},indent=2)+'\n')
    print(json.dumps({'docx':str(target),'appendix_pages':len(appendix),'equation_blocks':len(FORMULAS)},indent=2))


if __name__=='__main__':
    main()
