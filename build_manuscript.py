"""Build the evidence-linked research manuscript after the complete benchmark."""
from pathlib import Path
from copy import deepcopy
import ast
import io
import json
import math
import os
import subprocess
import tokenize
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd
from native_equations import prepare_equations

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'
EQ=ROOT/'tmp'/'equations'
RESULTS=ROOT/'results'

EQUATIONS={
 'model':r'\begin{aligned}V(c)&=\max_{x\geq0,\ Ax\leq c} U(x),\\U(x)&=\sum_{i=1}^{d}w_i\log(x_i+\xi_i),\qquad W=\sum_i w_i.\end{aligned}',
 'conjugate':r'\phi_i(q)=\begin{cases}w_i[\log(w_i/q)-1]+\xi_iq,&q\xi_i<w_i,\\w_i\log\xi_i,&q\xi_i\geq w_i,\end{cases}\quad q>0.',
 'dual':r'q=A^\top\lambda>0,\qquad B(c,\lambda)=c^\top\lambda+\sum_i\phi_i(q_i).',
 'certificate':r'0\leq V(c^\prime)-U(x)\leq B(c^\prime,\lambda)-U(x),\quad x\geq0,\ Ax\leq c^\prime,\ \lambda\geq0.',
 'geometric':r'\prod_i\left(\frac{x_i+\xi_i}{x_i^\star+\xi_i}\right)^{w_i/W}\geq e^{-\epsilon}.',
 'ray':r'\begin{aligned}h(t)&=B(c^\prime,t\lambda_0),\quad \gamma={c^\prime}^{\!\top}\lambda_0,\\h^\prime(t)&=\gamma-\sum_i(w_i/t-\xi_iq_{0i})_+.\end{aligned}',
 'waterfill':r'\gamma=\sum_iw_i(b-\alpha_i)_+,\qquad b=1/t,\quad\alpha_i=\xi_iq_{0i}/w_i.',
 'prefix':r'\Gamma_k=\alpha_kW_{k-1}-H_{k-1},\quad b=\frac{\gamma+H_k}{W_k},\quad H_k=\sum_{i\leq k}w_i\alpha_i.',
 'repair':r'\ell=Av,\quad r_e=\min\{1,c_e^\prime/\ell_e\},\quad x_i=v_i\min_{e:A_{ei}>0}r_e.',
 'theta':r'p_i=\frac{w_i}{W},\quad \theta_i=\frac{\xi_i}{x_{0i}+\xi_i},\quad z_i=s+(1-s)\theta_i,\quad\bar z=\mathbb E_p z_i.',
 'jensen':r'\begin{aligned}\frac{V(sc)-U(sx_0)}{W}&\leq J_s=\log\bar z-\mathbb E_p\log z_i\\J_s&\leq\frac{(1-s)^2}{2s^2}\operatorname{Var}_p(\theta),\qquad 0<s\leq1.\end{aligned}',
 'jensenproof':r'B_{\rm rel}(sc,t\lambda_0)=tW\bar z+U(x_0)-W-W\log t.',
 'smallchange':r'\frac{\delta^2}{2}\operatorname{Var}_p(\theta)\leq J_{1-\delta}\leq\frac{\delta^2}{2(1-\delta)^2}\operatorname{Var}_p(\theta).',
 'threshold':r'\delta\leq\frac{a}{1+a},\qquad a=\sqrt{\frac{2\epsilon}{\operatorname{Var}_p(\theta)}}.',
 'finitepars':r'\begin{aligned}r_i&=\frac{q_{0i}(x_{0i}+\xi_i)}{w_i},\quad \sigma=\frac{\lambda_0^\top(c-Ax_0)}{W},\\a_i&=r_i[s+(1-s)\theta_i],\qquad\bar a=\mathbb E_p a_i.\end{aligned}',
 'finitebound':r'\begin{aligned}\frac{V(sc)-U(sx_0)}{W}&\leq\log(\bar a+s\sigma)-\mathbb E_p\log a_i\\\log(\bar a+s\sigma)-\mathbb E_p\log a_i&=\log\bar a-\mathbb E_p\log a_i+\log(1+s\sigma/\bar a).\end{aligned}',
 'finitevar':r'\log(\bar a+s\sigma)-\mathbb E_p\log a_i\leq\frac{\operatorname{Var}_p(a)}{2(\min_i a_i)^2}+\log(1+s\sigma/\bar a).',
 'finiteproof':r'D=sc^\top\lambda_0+\sum_i\xi_iq_{0i}=W(\bar a+s\sigma),\qquad t^\star=W/D.',
 'alphamodel':r'U_\alpha(x)=\sum_i\frac{w_i(x_i+\xi_i)^{1-\alpha}}{1-\alpha},\qquad \alpha>0,\ \alpha\neq1.',
 'alphapars':r'v_i=w_i(x_{0i}+\xi_i)^{1-\alpha},\quad M=\sum_i v_i,\quad p_i=v_i/M.',
 'alphabound':r'\begin{aligned}\frac{V_\alpha(sc)-U_\alpha(sx_0)}{M}&\leq\frac{\bar z^{1-\alpha}-\mathbb E_p z_i^{1-\alpha}}{1-\alpha}\\\frac{\bar z^{1-\alpha}-\mathbb E_p z_i^{1-\alpha}}{1-\alpha}&\leq\frac{\alpha(1-s)^2}{2s^{\alpha+1}}\operatorname{Var}_p(\theta),\quad 0<s\leq1.\end{aligned}',
 'alpharay':r'\begin{gathered}a_i=(w_i/q_{0i})^{1/\alpha},\quad g_i=q_{0i}a_i,\quad\beta_i=\xi_i/a_i,\\\gamma=\sum_i g_i(b-\beta_i)_+,\qquad t=b^{-\alpha}.\end{gathered}',
 'sharp':r'\xi_i=w_i\theta_i,\quad x_{0i}=w_i(1-\theta_i),\quad c=\sum_i x_{0i},\quad\lambda_0=1.',
 'counter':r'x^\star=(r,1),\quad t^\star=\frac{2}{1+r},\quad\frac{\min_tB((r,1),t(1,1))-U(x^\star)}{2}=\frac12\log\frac{(1+r)^2}{4r}.',
 'decision':r'\Delta_{\rm ray}=\frac{\min_{t>0}B(c^\prime,t\lambda_0)-\max\{U(x_{\rm anchor}),U(x_{\rm ray})\}}{W}.',
 'references':r'\frac{L_{\rm ref}-U(x)}{W}\leq\frac{V(c^\prime)-U(x)}{W}\leq\frac{B_{\rm ref}-U(x)}{W}.',
 'profile':r'\rho_{p,m}=T_{p,m}/\min_jT_{p,j},\qquad P_m(\tau)=\frac{\#\{p:\rho_{p,m}\leq\tau\}}{\#\{p\}}.',
}


def field(paragraph,instruction):
    run=paragraph.add_run()
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=instruction
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,text,end])


def hyperlink(paragraph,label,url):
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');prop=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'195D78');prop.append(color);run.append(prop)
    text=OxmlElement('w:t');text.set(qn('xml:space'),'preserve');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)


class Manuscript:
    def __init__(self,metadata):
        self.doc=Document();self.meta=metadata;self.transcript=[]
        for border in self.doc.styles._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        sec=self.doc.sections[0]
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68)
        sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
        sec.header_distance=Inches(.30);sec.footer_distance=Inches(.30)
        sec.different_first_page_header_footer=True
        for name in ['Normal','Body Text']:
            st=self.doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11)
            st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.line_spacing=1.12
            st.paragraph_format.space_after=Pt(5)
            st.paragraph_format.widow_control=True
        for name,size in [('Title',24),('Subtitle',11.5),('Heading 1',14),('Heading 2',11.5)]:
            st=self.doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size)
            st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 7)
            st.paragraph_format.space_after=Pt(6)
            st.paragraph_format.keep_with_next=True
            st.paragraph_format.keep_together=True
            st.font.bold=name!='Subtitle'
        for st in self.doc.styles:
            if st.type==1 and st._element.rPr is not None:
                fonts=st._element.rPr.find(qn('w:rFonts'))
                if fonts is not None:
                    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                        fonts.attrib.pop(qn('w:'+attr),None)
        self.doc.styles['Title'].paragraph_format.line_spacing=1.0
        self.doc.styles['Title'].paragraph_format.space_before=Pt(9)
        self.doc.styles['Title'].paragraph_format.space_after=Pt(9)
        self.doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.doc.styles['Subtitle'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.doc.styles['Subtitle'].font.italic=True
        self.doc.styles['Subtitle'].font.color.rgb=RGBColor.from_string('404040')
        st=self.doc.styles['Caption'];st.font.name='Times New Roman';st.font.size=Pt(10)
        st.font.bold=False;st.font.italic=False
        st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.05
        st.paragraph_format.keep_together=True
        st.paragraph_format.keep_with_next=False
        hp=sec.header.paragraphs[0]
        hp.paragraph_format.space_after=Pt(0)
        rr=hp.add_run('Certified Reuse for Network Utility Maximisation')
        rr.font.name='Times New Roman';rr.font.size=Pt(8);rr.italic=True
        rr.font.color.rgb=RGBColor.from_string('606060')
        fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field(fp,'PAGE')
        for r in fp.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
        first=sec.first_page_footer.paragraphs[0];first.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field(first,'PAGE')
        for r in first.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
        props=self.doc.core_properties
        props.title='Certified Reuse for Network Utility Maximisation'
        props.author='Kai Chi'
        props.subject='Utility certificates, capacity changes, shifted logarithmic allocation'
        props.keywords='network utility maximisation; duality; price reuse; reproducible research'
        props.comments=''
        self.page_titles=[]

    def p(self,text,bold=False,size=None):
        previous=self.doc.paragraphs[-1] if self.doc.paragraphs else None
        p=self.doc.add_paragraph()
        p.paragraph_format.widow_control=True
        if bold:
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.keep_together=True
        if previous is not None and (previous.style.name=='Heading 1'
                or (previous.runs and all(r.bold for r in previous.runs if r.text))):
            p.paragraph_format.keep_together=True
        r=p.add_run(text);r.bold=bold
        if size:r.font.size=Pt(size)
        self.transcript.append(text+'\n')
        return p

    def heading(self,text,newpage=False):
        p=self.doc.add_paragraph(text,'Heading 1')
        if newpage:p.paragraph_format.page_break_before=True
        self.page_titles.append(text);self.transcript.append('\n# '+text+'\n')
        return p

    def sub(self,text):
        self.doc.add_paragraph(text,'Heading 2');self.transcript.append('\n## '+text+'\n')

    def eq(self,key):
        nodes=self.meta[key]
        if self.doc.paragraphs:
            self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
        for i,node in enumerate(nodes):
            p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5 if i==0 else 0)
            p.paragraph_format.space_after=Pt(8 if i==len(nodes)-1 else 2)
            p.paragraph_format.keep_with_next=i<len(nodes)-1
            p.paragraph_format.keep_together=True
            p._p.append(deepcopy(node))
        self.transcript.append('\n$$'+EQUATIONS[key]+'$$\n')

    def figure(self,name,caption,width=6.55):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'figures'/(name+'.png')),width=Inches(width))
        cp=self.doc.add_paragraph(style='Caption')
        label,body=caption.split('. ',1)
        cp.add_run(label+'. ').bold=True
        cp.add_run(body)
        self.transcript.append('\n![Scientific figure](../figures/'+name+'.png)\n'+caption+'\n')

    def table(self,headers,rows,widths,font=9.5):
        table=self.doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        pr=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for name in ['top','left','bottom','right','insideH','insideV']:
            tag=OxmlElement('w:'+name);tag.set(qn('w:val'),'single');tag.set(qn('w:sz'),'5');tag.set(qn('w:color'),'D9D9D9');borders.append(tag)
        pr.append(borders)
        for col,w in zip(table.columns,widths):col.width=Inches(w)
        for j,h in enumerate(headers):table.rows[0].cells[j].text=str(h)
        for row in rows:
            cells=table.add_row().cells
            for j,v in enumerate(row):cells[j].text=str(v)
        for i,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if i==0:
                repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                for name,value in [('top','80'),('bottom','80'),('left','95'),('right','95')]:
                    el=OxmlElement('w:'+name);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');marg.append(el)
                cp.append(marg)
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF1');cp.append(sh)
                elif i%2==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F7F8FA');cp.append(sh)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.line_spacing=1.04
                    p.paragraph_format.keep_with_next=i<len(table.rows)-1
                    p.paragraph_format.keep_together=True
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name='Arial';r.font.size=Pt(font);r.bold=i==0
                        r.font.color.rgb=RGBColor(0,0,0)
        gap=self.doc.add_paragraph()
        gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.space_after=Pt(4)
        gap.paragraph_format.line_spacing=Pt(1)
        gap.add_run().font.size=Pt(1)
        self.transcript.append('| '+' | '.join(headers)+' |\n')
        self.transcript.append('| '+' | '.join(['---']*len(headers))+' |\n')
        for row in rows:self.transcript.append('| '+' | '.join(map(str,row))+' |\n')

    def reference(self,number,text,url):
        p=self.p('['+str(number)+'] '+text,size=10)
        p.paragraph_format.space_after=Pt(7)
        p.paragraph_format.line_spacing=1.08
        p.paragraph_format.left_indent=Inches(.24)
        p.paragraph_format.first_line_indent=Inches(-.24)
        p.paragraph_format.keep_together=True
        hyperlink(p,'  Source',url)
        self.transcript.append(url+'\n')

    def code(self,source):
        for line in source.splitlines():
            p=self.doc.add_paragraph()
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(11)
            p.paragraph_format.widow_control=False
            r=p.add_run(line if line else ' ');r.font.name='Courier New';r.font.size=Pt(8.7)
        self.transcript.append('\n~~~python\n'+source+'\n~~~\n')


def wrap_code(source,width=88):
    """Wrap within brackets at safe token boundaries and preserve the AST."""
    out=[]
    for line in source.splitlines():
        if len(line)<=width:
            out.append(line);continue
        indent=len(line)-len(line.lstrip())
        if ' = ' in line and not line.lstrip().startswith(('if ','while ','def ')):
            left,right=line.split(' = ',1)
            line=left+' = ('+right+')'
        try:
            tokens=list(tokenize.generate_tokens(io.StringIO(line.lstrip()).readline))
        except tokenize.TokenError:
            out.append(line);continue
        depth=0;points=[]
        for tok in tokens:
            if tok.string in ['(', '[', '{']:depth+=1
            if tok.string in [',','+','-','*','/','**'] and depth>0:
                points.append(tok.end[1]+indent)
            if tok.string in [')',']','}']:depth-=1
        start=0;continuation=' '*(indent+4)
        while len(line)-start+(0 if start==0 else len(continuation))>width:
            limit=start+width-(0 if start==0 else len(continuation))
            candidates=[v for v in points if start<v<=limit]
            if not candidates:break
            cut=max(candidates)
            piece=line[start:cut].strip() if start else line[start:cut].rstrip()
            out.append((continuation if start else '')+piece)
            start=cut
            while start<len(line) and line[start]==' ':start+=1
        out.append((continuation if start else '')+line[start:].rstrip())
    result='\n'.join(out)+'\n'
    assert ast.dump(ast.parse(source),include_attributes=False)==ast.dump(ast.parse(result),include_attributes=False)
    return result


def appendix_sources():
    tree=ast.parse((ROOT/'certified_num.py').read_text())
    nodes={n.name:n for n in tree.body if isinstance(n,(ast.ClassDef,ast.FunctionDef))}
    chosen=[]
    # Full selected functions or classes; presentation wrapping preserves AST.
    model=nodes['NUM']
    selected={'utility','conjugate','rates','dual','repair'}
    body=[n for n in model.body if isinstance(n,ast.FunctionDef) and n.name in selected]
    excerpt=ast.ClassDef(name='NUM',bases=[],keywords=[],body=body,decorator_list=[])
    chosen.append(('Dual evaluation and feasibility repair',ast.unparse(ast.fix_missing_locations(excerpt))))
    ray=nodes['PriceRay']
    ray_excerpt=ast.ClassDef(name='PriceRay',bases=[],keywords=[],body=ray.body,decorator_list=ray.decorator_list)
    chosen.append(('Cached price ray',ast.unparse(ast.fix_missing_locations(ray_excerpt))))
    chosen.append(('Common stopping rule',ast.unparse(nodes['solve'])))
    chosen.append(('Measured method wrapper',ast.unparse(nodes['run_method'])))
    return [(title,wrap_code(code)) for title,code in chosen]


def main():
    s=json.loads((RESULTS/'summary.json').read_text())
    env=json.loads((RESULTS/'environment.json').read_text())
    audit=json.loads((RESULTS/'saved_result_audit.json').read_text())
    assert audit['passed']
    ablation=json.loads((RESULTS/'certificate_ablation_summary.json').read_text())
    finite=json.loads((RESULTS/'finite_anchor_validation.json').read_text())
    alpha=json.loads((RESULTS/'alpha_extension_validation.json').read_text())
    primal=json.loads((RESULTS/'independent_primal_checks.json').read_text())
    theory=json.loads((RESULTS/'theory_validation.json').read_text())
    cases=pd.read_csv(RESULTS/'cases.csv')
    anchors=pd.read_csv(RESULTS/'anchors.csv')
    core=cases[cases.split=='core']
    topologies=json.loads((RESULTS/'topologies.json').read_text())
    OUT.mkdir(exist_ok=True);EQ.mkdir(parents=True,exist_ok=True)
    m=Manuscript(prepare_equations(EQUATIONS,EQ))
    pair=s['paired_timing'];tolrows=[z for z in s['tolerance_sensitivity'] if z['tolerance']==1e-3]
    good=sum(z['reference_good'] for z in tolrows)
    good_rejected=sum(z['reference_good_rejected'] for z in tolrows)
    labels={'cold':'Cold prices','warm':'Warm prices','ray_warm':'Ray warm start','reuse':'Certified reuse'}
    shocklabels={'uniform_10':'Uniform 10%','uniform_30':'Uniform 30%','local_30':'Local 30%','heterogeneous_30':'Heterogeneous 30%'}
    acceptance=s['reuse_acceptances_core']
    uniform_acceptance=int(core.loc[core.shock.str.startswith('uniform'),'accept_0.001'].sum())
    def ratio(method):
        z=pair[method]
        return f"{z['median_time_ratio']:.2f} [{z['ratio_cluster_ci'][0]:.2f}, {z['ratio_cluster_ci'][1]:.2f}]"
    def median_ms(values):return f'{1000*values.median():.2f}'

    m.doc.add_paragraph('Certified Reuse for\nNetwork Utility Maximisation','Title')
    m.doc.add_paragraph('Utility error bounds after capacity changes','Subtitle')
    m.transcript.append('# Certified Reuse for Network Utility Maximisation\n\nUtility error bounds after capacity changes\n')
    author=m.p('Kai Chi\n22 October 2026',size=10.5)
    author.alignment=WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_before=Pt(4);author.paragraph_format.space_after=Pt(12)
    m.sub('Abstract')
    m.p(f"Repeated resource allocation creates a decision problem beyond solving each new optimum: can an inexpensive update be certified as sufficiently good? This study considers weighted shifted-log network utility maximisation with fixed routes and hard link capacities. A cached dual price direction gives an exact scalar water-filling problem and a valid upper bound for any feasible candidate. A sharp weighted Jensen bound explains when uniformly scaling a nominal optimum has small utility regret; a finite-anchor extension allows approximate prices and zero nominal rates. A two-link counterexample shows that the same price ray can yield an unboundedly loose certificate even when feasibility repair is exactly optimal. In a prespecified CPU study on six public topology structures, {acceptance} of 576 capacity cases were certified without iterative optimisation at a gap per total utility weight of 0.001. The median reuse-to-ray-warm runtime ratio was {pair['ray_warm']['median_time_ratio']:.2f} on {pair['ray_warm']['common_solved_cases']} cases solved by both methods. The study retains all reference and measured failures, includes 24 scaling cases up to 100,000 flows, and separates synthetic workloads from public graph data. For nonuniform changes, the median ratio was {s['stratified_pairing']['nonuniform']['ray_warm']['median_time_ratio']:.2f}; the precheck did not improve median runtime in that regime. The evidence distinguishes a valid utility guarantee from a broadly useful computational shortcut.")
    m.sub('Research question')
    m.p("After a capacity reduction, a repaired allocation may already have acceptable utility. Reoptimising by default spends computation before establishing that it is necessary. Reusing by default leaves its loss unknown. The useful question is whether a cheap, auditable bound can separate the two cases, and whether computing that bound costs less than a strong direct update.")
    m.p("The analysis therefore addresses three distinct claims: validity of an acceptance decision, coverage of that decision, and its computational value. Validity follows from weak duality. Coverage depends on the available dual directions and candidate quality. Computational value must include failed prechecks, preparation and a competitive warm start. These claims require different evidence.")

    m.heading('1 Relationship to existing research',newpage=True)
    m.p("The starting paper by Tian et al., with Jiahai Yang among its authors, studies centralised NUM using a smooth penalty and accelerated gradient iterations [1]. The present problem has hard capacities and repeated capacity changes. It does not reproduce that method or transfer its convergence theorem. A comparison against plain gradient descent would not establish a competitive new solver.")
    m.p("Lagrangian bounds, coordinate conjugates and KKT conditions are standard convex optimisation [2]. Uniform scaling for unshifted logarithmic allocation is a known iso-elasticity property [3]. Sorted water filling has an extensive algorithmic literature [4]. Dual scaling and gap-based safe decisions also appear in other optimisation settings [9]. These are building blocks, not discoveries claimed by this study.")
    m.p("The closest adaptation reference is Hou's zero-shot Lagrangian update [6], which uses sensitivity-based multiplier adjustment under different information and resource assumptions. Reusing prices after a capacity change is consequently not a novelty claim. The narrower candidate contribution here is the stated shifted-log regret bound, its inexact-anchor decomposition, and an explicit account of when a single-direction certificate fails. Priority for those statements remains unverified.")
    m.table(['Reference family','Established scope','Implication here'],[
        ('Packing proportional fairness [5]','Accelerated packing algorithms','No best-solver claim'),
        ('Dynamic multiplier update [6]','Capacity adaptation','Reuse alone is not new'),
        ('GPU proximal messaging [7]','Large NUM and warm starts','No matched GPU comparison'),
        ('GATE traffic engineering [8]','Multipath and production traces','Richer operational model'),
    ],[2.05,2.55,2.25],font=9.3)
    m.p("Recent GPU NUM and WAN traffic-engineering systems [7,8] make large-scale and warm-start claims insufficient by themselves. Their models, hardware and datasets differ from this CPU benchmark. Published wall-clock numbers are not imported into a misleading speedup table. The literature search is targeted and current to the retrieval date, rather than a proof that no equivalent result exists.")

    m.heading('2 Model and a valid utility certificate')
    m.p("Let A be a nonnegative link-route matrix with m rows and d nonempty columns. Capacities are strictly positive, weights are positive, and utility shifts are nonnegative. Routes and utility parameters stay fixed while capacities change. Natural logarithms are used.")
    m.eq('model')
    m.p("Every flow is bounded by a capacity on its route. A positive feasible allocation exists, and strict concavity gives a unique optimal allocation. When a shift is zero, finite utility requires a positive rate for that flow. Positive shifts permit optimal zero rates.")
    m.p('Proposition 1  Feasible primal and dual values certify regret',bold=True)
    m.p("For nonnegative prices with positive route prices, the scalar maximiser is the positive part of the weight divided by its route price, minus its shift. The coordinate conjugate and the dual upper bound are")
    m.eq('conjugate');m.eq('dual')
    m.eq('certificate')
    m.p("Proof. Add the nonnegative term given by prices times unused capacity to any feasible utility. Maximising the resulting Lagrangian over the nonnegative rates separates by coordinate and gives B. Maximising the original utility over the feasible region proves the inequality. No optimal-anchor assumption is needed.")
    m.p("The stopping target divides this gap by W. It bounds aggregate weighted log utility, not the maximum loss on an individual route. Equivalently, for an accepted allocation and the updated optimum,")
    m.eq('geometric')
    m.p("At the primary tolerance 0.001 the weighted geometric mean ratio of shifted rates is at least about 0.999. A particular flow can still fare much worse. The implementation uses float64 feasibility checks and a small downward rounding guard; it is not an interval-arithmetic proof.")

    m.heading('3 Exact optimisation on a cached price ray')
    m.p("Fix any nonnegative nominal price vector whose route prices are positive. Restrict updated prices to positive scalar multiples of this vector. This gives the strongest bound on that ray, not necessarily the strongest bound in the full dual space.")
    m.eq('ray')
    m.p('Proposition 2  The scalar minimiser is an active set water filling solution',bold=True)
    m.eq('waterfill')
    m.p("Proof. The derivative follows by differentiating each coordinate supremum at its unique maximiser. Its subtracted sum decreases continuously from infinity to zero. Since gamma is positive, the derivative has exactly one zero, which minimises the convex ray objective. The positive part is essential when some shifted-log rates vanish.")
    m.p("Sort the alpha values. Cache prefix weight sums and weighted alpha sums. The nondecreasing breakpoints identify the active prefix by binary search; the associated inverse multiplier is")
    m.eq('prefix')
    m.p("Indeed, successive breakpoints differ by the next alpha spacing times the current prefix weight, so ties do not invalidate the search. If an active coordinate lies exactly at its breakpoint it contributes zero. The code handles this through a consistent right-sided search.")
    m.table(['Operation','Cost'],[
        ('Prepare route prices and sorted prefix cache','O(nnz(A) + d log d)'),
        ('Compute gamma for a new capacity vector','O(m)'),
        ('Query the cached scalar active prefix','O(log d)'),
        ('Construct rates, repair and certify','O(nnz(A) + d + m)'),
    ],[4.6,2.25])
    m.p("Only the scalar prefix query is logarithmic time. The full online update still touches routes and flows. This is an application of standard water-filling machinery; an exact ray query is not presented as an independently established new optimisation algorithm.")

    m.heading('4 A sharp bound for uniform capacity changes')
    m.p('Proposition 3  Shift heterogeneity controls uniform scaling regret',bold=True)
    m.p("Assume the nominal allocation is an exact optimum with every nominal rate positive, and nominal prices satisfy KKT. For a uniform contraction factor s, define the shift fraction and its weighted transformed mean by")
    m.eq('theta');m.eq('jensen')
    m.p("Proof. Stationarity gives each route price as the weight divided by the nominal shifted rate. Complementary slackness gives the nominal capacity-price product as W times one minus the weighted mean shift fraction. The shift-price product is W times that mean. Relaxing each conjugate from nonnegative rates to rates above minus the shift yields")
    m.eq('jensenproof')
    m.p("Its minimiser is the reciprocal of the mean z. The resulting upper bound is nominal utility plus W times the log of that mean. Scaling the primal allocation changes utility by the weighted sum of log z. Subtraction proves the Jensen bound. Taylor's inequality for minus log, whose second derivative is at most one divided by s squared on this interval, proves the variance bound.")
    m.p("If all shifts vanish, or all shift fractions coincide, the bound is zero. The first case is known proportional-fair scaling [3]. The second shows why variation of relative shifts, rather than the mere presence of a shift or a particular graph, is the relevant quantity.")
    m.eq('smallchange')
    m.p("Here delta is one minus s. These two inequalities bound the Jensen expression itself; only the upper one bounds regret on a general network. The leading term is one half of delta squared times shift-fraction variance. A sufficient contraction limit for a prescribed utility tolerance is")
    m.eq('threshold')
    m.p("The exact KKT and positive-rate assumptions are not silently assigned to numerical anchors. The following extension supplies a computable replacement when those assumptions are unavailable.")

    m.heading('5 A bound for inexact nominal solutions')
    m.p('Proposition 4  Stationarity variation and unused capacity give explicit correction terms',bold=True)
    m.p("Let the nominal allocation be any feasible finite-utility point, and let nominal prices be any nonnegative vector with positive route prices. Nominal zero rates are allowed when their shifts are positive. For every uniform factor s greater than zero, define")
    m.eq('finitepars');m.eq('finitebound')
    m.p("Proof. The scaled allocation is feasible. Relax each scalar conjugate to rates greater than minus its shift, as in Proposition 3. The coefficient multiplying the price scale is")
    m.eq('finiteproof')
    m.p("Subtracting the utility of the scaled allocation from that relaxed dual gives, after division by W, tD/W minus one minus log t minus the weighted mean log a. Minimising over t gives the displayed bound. Since the true scalar conjugate is no larger than the relaxed one, the exact ray certificate for this same candidate is at least as strong.")
    m.eq('finitevar')
    m.p("The last inequality follows by Taylor's theorem on the positive range of a. At an exact positive-rate KKT anchor, every stationarity ratio r equals one and sigma vanishes. The expression reduces to Proposition 3. For capacity expansions the same derivation holds; the exact-anchor variance denominator becomes the square of the smaller of s and one.")
    m.p("This decomposition exposes two issues hidden by exactness assumptions. Nonstationary prices change the spread of a, while price-weighted unused capacity adds a nonnegative slack term. The Jensen term is also the KL divergence from p to the distribution proportional to p times a. This is an interpretation of the bound, not a new entropy identity. Both terms are computable from the stored allocation, routes and prices.")
    m.p(f"The supplement was derived while the fixed benchmark ran and did not alter its methods or thresholds. In {finite['cases']} additional checks, including {finite['zero_nominal_rate_cases']} zero-rate anchors and {finite['capacity_expansions']} expansions, the formula matched the explicitly evaluated relaxed dual to {finite['maximum_formula_identity_error']:.2g}; the exact ray bound never exceeded it. These checks support algebraic consistency, not publication priority.")

    m.heading('6 Extension to shifted alpha fair utility')
    m.p("The argument extends beyond logarithmic utility. For positive alpha other than one, use the shifted alpha-fair utility below. Its limit at one is weighted shifted-log utility after removal of an allocation-independent constant.")
    m.eq('alphamodel')
    m.p('Proposition 5  A weighted power Jensen expression controls scaling regret',bold=True)
    m.p("Under the exact positive-rate nominal KKT assumptions of Proposition 3, retain the same shift fractions and z values, but change the weights used in their expectation to")
    m.eq('alphapars');m.eq('alphabound')
    m.p("Proof. Nominal stationarity gives route prices equal to weight times the nominal shifted rate to the power minus alpha. The relaxed ray dual is M times alpha divided by one minus alpha times t to the power (alpha minus one) divided by alpha, plus t times M times the mean z. Its minimiser is the mean z to the power minus alpha. Subtracting the scaled primal utility gives the power Jensen expression. The second derivative of the utility of z is minus alpha times z to the power minus alpha minus one, giving the variance bound.")
    m.p("For alpha above one, both the power Jensen numerator and its denominator are nonpositive; their ratio remains nonnegative. A one-link construction with weight equal to the nominal shifted rate to the power alpha makes the first bound sharp when updated rates remain positive. Vanishing shifts recover known alpha-fair iso-elasticity [3].")
    m.p("The exact nonnegative price-ray minimisation also remains a weighted water-filling problem. For arbitrary valid nominal prices, define")
    m.eq('alpharay')
    m.p("The scalar equation follows by differentiating the coordinate supremum, with rates equal to the positive part of a times b minus the shift. Thus the same sorted-prefix idea applies with different coefficients; the log implementation is its alpha-equals-one case.")
    m.p(f"The separate supplement verifies {alpha['independent_ray_cases']} independent scalar comparisons and {alpha['sharpness_cases']} sharpness constructions across alpha values 0.2, 0.5, 1, 2 and 4, plus continuity near one. The six-topology timing experiment remains logarithmic only. No measured runtime conclusion is transferred to other alpha values.")

    m.heading('7 Sharpness and a structural limitation')
    m.figure('01_theory_and_counterexample','Figure 1. Left: a one-link family realises the Jensen bound as exact regret; the variance expression is an upper bound. Right: a two-link repaired allocation is optimal while its best fixed-ray certificate diverges.')
    m.p("Sharpness. On one link choose positive weights and shift fractions strictly between zero and one, and construct")
    m.eq('sharp')
    m.p("Nominal prices equal one and the stated allocation satisfies KKT. If the largest shift fraction is below the transformed mean z, all updated optimal rates remain positive. The single price ray is then the full dual, and the first inequality in Proposition 3 is equality. Nine saved examples with two, five and forty flows agree within 1e-12.")
    m.p('Proposition 6  An optimal repaired allocation can have an arbitrarily loose ray certificate',bold=True)
    m.p("Take two disjoint links, identity routing, unit weights, zero shifts and nominal capacities and prices both equal to (1,1). Reduce only the first capacity to r. Routewise repair gives (r,1), which is exactly optimal, but")
    m.eq('counter')
    m.p("Proof. The ray dual minimiser is two divided by one plus r. Its dual value is twice the log of (1+r)/2, whereas the optimal primal utility is log r. Their difference gives the expression, which diverges as r tends to zero.")
    m.p("A failed certificate therefore does not establish a poor allocation. The ray can accept safely but miss excellent candidates. No uniform tightness claim under arbitrary heterogeneous changes follows from a successful uniform-scaling example.")

    m.heading('8 Feasible candidates and measured decision rules')
    m.p("A candidate must be feasible before weak duality yields a useful regret certificate. For any nonnegative proposed rates, compute the link loads and downscale each flow by the smallest admissible factor on its route. Set a zero-load link factor to one.")
    m.eq('repair')
    m.p("Every term using a link is multiplied by at most that link's factor, so the updated load is at most its capacity. This is a one-pass feasibility repair, not a Euclidean projection or a fairness-optimal repair. It works for the nonnegative routing model and is not claimed as new.")
    m.p("The reusable method repairs both the nominal allocation and the allocation stationary at the best ray prices, then keeps the higher utility. Its decision quantity is")
    m.eq('decision')
    m.p("If the quantity is at most the target, it returns the feasible candidate. Otherwise it invokes the common solver from the ray prices with that candidate as a lower bound. Every precheck and fallback is timed.")
    m.table(['Measured method','Price initialisation','Initial primal information'],[
        ('Cold prices','Constant 0.2','Repaired nominal allocation'),
        ('Warm prices','Nominal prices','Repaired nominal allocation'),
        ('Ray warm start','Exact scalar ray minimiser','Repaired nominal allocation'),
        ('Certified reuse','Ray check then conditional solve','Better of two repaired candidates'),
    ],[1.45,2.35,3.05],font=9.3)
    m.p("All methods use the same L-BFGS-B implementation and the same feasible-primal/dual target at every function-gradient evaluation. The cold-price baseline still receives nominal primal information and is not a completely cold solve. Tight references alone may use safeguarded Newton refinement. Price lower bounds remain valid nonnegative dual points and do not relax the stopping target.")
    m.p("A separate two-by-two certificate ablation holds the primal candidate fixed while changing prices, and holds prices fixed while changing the candidate. The simpler fixed-price versus full-ray comparison changes both components and cannot isolate the scalar price effect by itself.")

    m.heading('9 Experimental design and public topology data')
    m.p("Six graph structures come from Internet Topology Zoo [10], pinned to a repository commit with exact hashes and the CC BY 4.0 licence. Self loops are removed, parallel edges collapsed, the largest component retained and edges expanded into directed links. Deterministic shortest paths connect randomly sampled endpoints. Historical graph maps are not present-day traffic measurements.")
    m.table(['Topology','Used nodes','Undirected edges'],[
        (name,topologies[name]['used_nodes'],topologies[name]['simple_undirected_edges'])
        for name in ['Abilene','Geant2012','Cernet','Colt','UsCarrier','Cogentco']
    ],[3.8,1.45,1.6])
    m.p("For each topology, the core design crosses three seeds, 1,000 and 10,000 flows, two capacity laws and shifts of zero and 0.5. Weights are lognormal and normalised to mean one. The load-coupled law multiplies link offered load by 0.8 and an independent uniform factor between 0.7 and 1.3. The independent law multiplies mean offered load by 0.8 and an independent lognormal factor with log standard deviation 0.7.")
    m.p("Each of 144 nominal anchors produces four shocks: uniform reductions of 10% and 30%; a 30% reduction on a random 5% of used directed links, rounded up; and independent reductions between zero and 30% across links. This gives 576 core cases sharing 36 route matrices. A prespecified extension adds 24 cases on Abilene and Cogentco at 1,000, 10,000 and 100,000 flows. Capacities, endpoints, weights and shocks are all generated.")
    m.table(['Protocol element','Specified value'],[
        ('Primary gap per total weight','0.001'),
        ('Nominal and updated reference target','0.0000001'),
        ('Precheck tolerance sensitivity','0.0001, 0.001, 0.01'),
        ('Timing repetitions and methods','3 repetitions, 4 methods'),
        ('Core cases and scaling cases','576 and 24'),
        ('Iteration and function-evaluation caps','3000 and 45000'),
    ],[4.15,2.7])
    m.p("Method order rotates across cases and repetitions. Timings include the online solver, repair and certification. Cache preparation is recorded and amortised over four core or two scaling updates. A second analysis adds the nominal anchor solve. The design uses fixed nominal anchors, not a sequential controller or a measured failure trace. All settings were fixed before main execution; this was not an external preregistration.")

    m.heading('10 Certificate coverage and controlled ablation')
    m.figure('02_certificate_coverage','Figure 2. Fraction of core cases accepted at a 0.001 gap per total weight. Each heatmap cell contains 36 cases. The left method uses unchanged prices and the repaired anchor; the right uses the best ray and the better candidate, so both components change.')
    m.p(f"The full reuse precheck accepted {acceptance}/576 core cases ({100*acceptance/576:.1f}%). The unchanged-price repaired-anchor check accepted {s['fixed_acceptances_core']}/576. The {s['known_scaling_cases']} unshifted uniform-scaling cases are a known-property sanity check; {s['known_scaling_acceptances']} were accepted. Among all accepted cases, {uniform_acceptance} came from uniform reductions and {acceptance-uniform_acceptance} from local or heterogeneous reductions. The remaining {acceptance-s['known_scaling_acceptances']} acceptances beyond known unshifted scaling had positive utility shifts.")
    m.table(['Shock type','Full reuse accepted','Core cases'],[
        (shocklabels[z['shock']],z['accepted'],z['cases']) for z in tolrows
    ],[3.65,1.65,1.55])
    acc=ablation['acceptances_at_0.001']
    m.table(['Certificate prices','Repaired anchor','Better candidate'],[
        ('Unchanged nominal prices',acc['fixed_anchor'],acc['fixed_best']),
        ('Best scalar price ray',acc['ray_anchor'],acc['ray_best']),
    ],[3.65,1.65,1.55])
    m.p("The two-by-two table reports acceptances over the same 576 cases. Moving down a column isolates the stronger dual point; moving across a row isolates the primal candidate. The full combination cannot be credited entirely to either component. These untimed controls were specified during main execution before comparative results were inspected.")
    m.p("Coverage depends on the distribution of shocks, utility shifts and capacity laws. The six topology families were selected for transparent diversity rather than sampled from a population of operators, so the acceptance fraction is a result for this benchmark.")

    m.heading('11 A rejected certificate can hide a good allocation')
    m.figure('03_certificate_conservatism','Figure 3. Cached ray gap against the candidate regret upper bound from a tight updated reference. Dashed lines mark 0.001. Points left of the vertical line but above the horizontal line are reference-good candidates that the ray does not certify. Plot values are floored at 1e-10 for display.')
    m.eq('references')
    m.p(f"The reference interval identifies {good}/576 core candidates as within tolerance using its upper endpoint. Of these, {good_rejected} were rejected by the cached-ray certificate. This is a conditional missed-certification count, not an estimate of unsafe acceptance. Cases whose reference interval straddles the tolerance remain unresolved rather than being forced into a binary label.")
    m.p("A reference solve supplies both a feasible lower bound and a dual upper bound. Neither endpoint is silently called the exact optimum. If the reference lower bound lies below the reuse candidate, the regret lower endpoint can be truncated to zero. The upper endpoint remains valid even if a reference fails its target.")
    m.p("The two-link counterexample explains why missed certification is structurally possible: capacity changes may require rotation of the price vector, not just a scale change. Better primal repair alone cannot guarantee a tight certificate if the available dual direction is inadequate.")
    m.p("This distinction changes the interpretation of the algorithm. A rejection is a request for stronger evidence or optimisation. It is not a diagnosis of excessive utility loss. An extension with a richer dual subspace might improve coverage, but would add preparation and online work; that extension has not been measured here.")
    m.p(f"The maximum updated reference gap was {s['maximum_reference_gap']:.3g}, with {s['reference_misses']} nominal or updated target misses recorded. Numerical evidence is supplementary to Proposition 1; it does not replace an analytical bound with an empirical safety rate.")

    m.heading('12 Runtime against common and strong warm starts')
    m.figure('04_runtime_profiles','Figure 4. Performance profiles [11] over 576 core cases. Left: online cost with amortised cache preparation. Right: the same cost plus the nominal anchor solve amortised over four updates. Unsolved methods receive infinite profile cost and remain in the denominator.')
    m.table(['Method','Solved cases','Median cost ms','Median solver FG calls'],[
        (labels[z['method']],f"{z['solved']}/{z['cases']}",f"{z['median_amortized_ms_all_budget']:.2f}",f"{z['median_fg_all']:.0f}")
        for z in s['method_table']
    ],[2.3,1.25,1.65,1.65],font=9.3)
    m.p("Median cost in this table includes capped-run execution time if any; it is not a conditional median over successes. A single low median can hide failures and case heterogeneity, so the performance profiles retain all cases and the paired comparisons below condition explicitly on both methods solving.")
    m.table(['Reuse divided by','Common solved','Median ratio and interval'],[
        (labels[k],pair[k]['common_solved_cases'],ratio(k)) for k in ['cold','warm','ray_warm']
    ],[2.3,1.25,3.3],font=9.3)
    m.p("Ratios below one favour reuse. Intervals are percentile bootstrap intervals for the median paired ratio, resampling six topology families with their settings intact. Three timing repetitions are first reduced to a per-case median. The intervals describe this limited graph collection; six clusters do not support strong population claims.")
    z=s['stratified_pairing']['nonuniform']['ray_warm']
    zz=s['stratified_pairing']['excluding_known_uniform']['ray_warm']
    m.p(f"Removing the known unshifted uniform-scaling cases gives a median reuse-to-ray-warm ratio of {zz['median_time_ratio']:.2f} over {zz['common_solved_cases']} cases solved by both methods. For local and heterogeneous shocks alone it is {z['median_time_ratio']:.2f} over {z['common_solved_cases']} cases. These subsets test whether the aggregate result is driven by an already tractable regime.")
    r=pair['ray_warm']['median_time_ratio']
    m.p(("The ray-warm baseline already captures much of the inexpensive price adjustment. " +
         ("The lower median reuse cost indicates an additional benefit on these cases solved by both methods, but not on every case. " if r<1 else
          "Reuse does not improve the median cost against that stronger baseline in these cases solved by both methods. ") +
         "Its fallback overhead remains part of the measured result. Neither the cold-price comparison nor skipping iterations alone establishes practical superiority."))

    m.heading('13 Scaling and the cost of reusable information')
    m.figure('05_scaling','Figure 5. Descriptive scaling extension on two topologies. Each flow count aggregates four settings, crossing two shifts and two shock types, at one seed. Lines are medians and shaded bands are interquartile ranges, not independent confidence intervals. All capped costs are retained.')
    m.p("The extension reaches 100,000 flows but has only 24 capacity cases and one seed. It probes execution growth rather than proving a scaling law or robustness to a distribution of large networks. Sparse route incidences, path lengths and the active bottleneck set all affect work beyond the number of flows. All four methods missed the target on the unshifted 100,000-flow Cogentco heterogeneous case in all three repetitions; those budget-limited times remain in the plot.")
    m.p(f"Across 156 anchors the median cache preparation cost was {s['cache_preparation_ms_median']:.2f} ms and the median nominal reference solve cost was {s['source_solve_ms_median']:.2f} ms. These are separate medians, not values that may be added to recover a median total. The original execution stopped after 154 complete anchors. The remaining two scaling anchors were resumed in a later process; all earlier within-case measurements were retained. Full original wall time is unavailable. The checkpoint and resume record are included. Completion-session CPU and resource limits are in results/hardware.json; processor metadata from the pre-resume segment were not recorded.")
    m.table(['Reuse divided by','Online ratio','Ratio with anchor cost'],[
        (labels[k],f"{pair[k]['median_time_ratio']:.2f}",f"{pair[k]['median_ratio_with_anchor']:.2f}") for k in ['cold','warm','ray_warm']
    ],[2.7,1.7,2.45])
    m.p("All four methods receive a nominal primal anchor. The second cost convention charges the same historical solve across methods and apportions it over the measured updates. It is a sensitivity analysis of accounting, not proof that four updates are a representative deployment horizon. With a single use, more preparation must be charged per decision; with many uses, it can be amortised more widely.")
    m.p("No GPU execution was performed. Recent NUM and traffic-engineering systems [7,8] cover different or much larger settings with different hardware. This extension is not evidence of parity or superiority to those systems. It also omits communication cost, solver integration and end-to-end controller latency.")
    m.p("A practical break-even calculation needs the actual frequency of reusable decisions, the cost of constructing and maintaining prices, and the latency target. Those operational observations are not supplied by the public graph structures.")

    m.heading('14 Tolerances failures and independent checks')
    m.figure('06_tolerance_sensitivity','Figure 6. Left: acceptance as the tolerance changes. Right: the fraction of reference-good candidates rejected, conditional on that reference-good set. Each shock type has 144 core cases; changing tolerance reuses the saved candidates rather than creating new networks.')
    m.p(f"No observed accepted case was contradicted by a reference lower bound at any of the three tested tolerances. There were {s['failed_measured_runs']} measured runs that missed the primary target. The full raw data and statuses are retained. This observation is not used to establish mathematical safety, which rests on a valid feasible-primal/dual bound.")
    m.table(['Check','Cases','Evidence'],[
        ('Independent scalar ray minimisation',str(theory['ray_trials']),f"Relative objective discrepancy {theory['max_ray_objective_relative_error']:.2g}"),
        ('Sharp one-link constructions','9','Jensen equality within 1e-12'),
        ('Independent primal SLSQP','4',f"Maximum utility discrepancy {max(abs(z['dual_vs_primal_gap_per_weight']) for z in primal):.2g}"),
        ('Finite-anchor extension','500',f"Formula discrepancy {finite['maximum_formula_identity_error']:.2g}"),
        ('Alpha-fair ray and sharpness','195',f"Ray relative discrepancy {alpha['maximum_relative_ray_error']:.2g}"),
        ('Saved problem recomputation','600',f"Maximum ray gap discrepancy {audit['maximum_reuse_gap_recomputation_error']:.2g}"),
    ],[2.55,.65,3.65],font=9.1)
    m.p("The independent primal checks use small 40-flow Abilene and Cernet problems with zero and positive shifts. They validate the dual implementation against a different optimisation formulation; they are not a modern competitive baseline. Directional derivatives, feasibility and weak duality are additionally checked.")
    m.p(f"Recomputing all saved arrays found maximum candidate overload {audit['maximum_reuse_overload']:.2g} relative to capacity. Recorded execution hashes and subsequent source text revisions are documented in the archive; the numerical operations are unchanged. The checks use ordinary floating-point arithmetic and do not cover every possible input, concurrency pattern or deployment failure.")

    m.heading('15 Interpretation limitations and research status')
    m.p("The strongest conclusion is conditional. A feasible allocation with a valid dual upper bound can be accepted without iteration; a sharp uniform-scaling bound explains one tractable regime. Restricting dual prices to one direction also creates a demonstrable blind spot. These statements can coexist: analytical validity does not imply high coverage, and high coverage does not guarantee lower runtime.")
    m.p("The specialised Jensen result gives a reasoned explanation for an observation that a small synthetic solver comparison could easily overstate. For unshifted logs, uniform scaling succeeds by a known structural property. With shifts, their relative dispersion controls a sharp upper bound. With inexact anchors, stationarity variation and price-weighted slack enter explicitly. This chain of reasoning is more informative than treating every capacity change as a generic optimisation task.")
    m.p("The evidence remains limited. Topology structures are public, but capacities and shocks are generated. Only six graph families and three core seeds are used. Routes are fixed, sessions can repeat endpoints, and demand bounds, multipath optimisation, arrivals, price privacy and an online sequence are absent. Aggregate utility is not a user-level service guarantee. Float64 checks do not provide machine-verified arithmetic.")
    m.p("Algorithmic competition is also incomplete. The shared L-BFGS-B implementation gives controlled comparisons of update choices, including a strong ray warm start; it does not establish dominance over current packing, conic or GPU solvers. Large-scale curves on two graphs do not close that gap. A claim about an operational deployment needs representative measurements or a carefully delimited theoretical application.")
    m.p("Priority is a separate unresolved question. Weak duality, water filling, feasible downscaling, uniform iso-elasticity and price reuse are established. The shifted-log bound, its sharpness family and finite-anchor decomposition are proved here as candidate specialised results. An unsuccessful keyword search is not a novelty proof. Domain-expert scrutiny against sensitivity and allocation theory is required.")
    m.p("Before a journal submission, the main decision is whether that specialised result is both new and substantial enough to carry a paper. If not, the research needs a demonstrably stronger certificate family or a clearly justified application with matched modern baselines. Independent proof and code review remain necessary. The present manuscript is research progress, not an assertion of Q1 readiness or acceptance.")
    m.sub('Reproducibility')
    m.p("The accompanying archive includes the full fixed protocol, all 600 capacity instances, 7,200 timing records, failures, source and data hashes, mathematical notes and figure scripts. It reproduces the numerical analysis with Python 3.12. The DOCX contains native editable Word equations, prose and tables. The appendix extracts come from the executed source with AST-preserving presentation wrapping.")

    m.heading('References',newpage=True)
    m.reference(1,"Y. Tian, Z. Wang, X. Yin, X. Shi, J. Yang and H. Zhang. Centralized Network Utility Maximization with Accelerated Gradient Method. IEEE ICNP, 2022. Author preprint arXiv:2408.08034v2, 2024.","https://arxiv.org/abs/2408.08034")
    m.reference(2,"S. Boyd and L. Vandenberghe. Convex Optimization. Cambridge University Press, 2004.","https://web.stanford.edu/~boyd/cvxbook/")
    m.reference(3,"S. Borst, N. Walton and B. Zwart. Network iso-elasticity and weighted alpha-fairness. arXiv:1201.2292, 2012.","https://arxiv.org/abs/1201.2292")
    m.reference(4,"D. P. Palomar and J. R. Fonollosa. Practical Algorithms for a Family of Waterfilling Solutions. IEEE Transactions on Signal Processing 53(2), 686-695, 2005. DOI 10.1109/TSP.2004.840816.","https://doi.org/10.1109/TSP.2004.840816")
    m.reference(5,"F. Criado, D. Martinez-Rubio and S. Pokutta. Fast Algorithms for Packing Proportional Fairness and its Dual. arXiv:2109.03678v2, 2022.","https://arxiv.org/abs/2109.03678")
    m.reference(6,"I.-H. Hou. Network Optimization in Dynamic Systems: Fast Adaptation via Zero-Shot Lagrangian Update. arXiv:2412.07865v1, 2024.","https://arxiv.org/abs/2412.07865")
    m.p("Reference versions are stated because the latest related preprints changed after their initial posting. A preprint is not labelled a peer-reviewed journal publication. For reference 4, bibliographic metadata and the author's publication list were verified; an access-blocked repository PDF was not represented as fully reviewed.",size=10)

    m.reference(7,"A. Sreekumar, A. Degleris and R. Rajagopal. Large-Scale Network Utility Maximization via GPU-Accelerated Proximal Message Passing. IEEE Control Systems Letters 9, 2807-2812, 2025. Related author version arXiv:2509.10722v2, revised September 2026.","https://arxiv.org/abs/2509.10722")
    m.reference(8,"Bothra et al. GATE: GPU-Accelerated Traffic Engineering for the WAN. arXiv:2605.01748v2, 29 September 2026.","https://arxiv.org/abs/2605.01748")
    m.reference(9,"E. Ndiaye, O. Fercoq, A. Gramfort and J. Salmon. Gap Safe Screening Rules for Sparsity Enforcing Penalties. Journal of Machine Learning Research 18(128), 1-33, 2017.","https://jmlr.org/papers/v18/16-577.html")
    m.reference(10,"S. Knight, H. X. Nguyen, N. Falkner, R. Bowden and M. Roughan. The Internet Topology Zoo. IEEE Journal on Selected Areas in Communications 29(9), 1765-1775, 2011. DOI 10.1109/JSAC.2011.111002.","https://doi.org/10.1109/JSAC.2011.111002")
    m.reference(11,"E. D. Dolan and J. J. Moré. Benchmarking Optimization Software with Performance Profiles. Mathematical Programming 91, 201-213, 2002. Author preprint arXiv:cs/0102001.","https://arxiv.org/abs/cs/0102001")
    m.sub('Data provenance')
    m.p("Graph data source: mroughan/InternetTopologyZoo, commit 51118849b99534f01010d885d66321433a7eef9e. The six source files and CC BY 4.0 licence are retained with hashes and source URLs. Only graph structures are observed; no real capacities, session traffic or capacity-loss traces are claimed.",size=10)

    extracts=appendix_sources()
    appendix_manifest=[]
    letter=ord('A')
    for label,code in extracts:
        lines=code.splitlines()
        # A complete node may span pages; no source lines are silently omitted.
        chunks=[lines[i:i+54] for i in range(0,len(lines),54)]
        for idx,chunk in enumerate(chunks):
            suffix=' continued' if idx else ''
            title=f'Appendix {chr(letter)} {label}{suffix}'
            m.heading(title,newpage=True)
            if idx==0:
                m.p("Extracted from certified_num.py. Presentation wrapping preserves the Python AST. The complete executable file and required surrounding definitions are supplied in the archive.",size=9.5)
            m.code('\n'.join(chunk))
            appendix_manifest.append({'heading':title,'lines':len(chunk),'maximum_line_length':max(map(len,chunk))})
        letter+=1
    target=OUT/'Certified_NUM_Research.docx'
    m.doc.save(target)
    (OUT/'Certified_NUM_Research.md').write_text('\n'.join(m.transcript))
    (RESULTS/'manuscript_build.json').write_text(json.dumps({'docx':str(target),'page_sections':m.page_titles,
        'appendix_extracts':appendix_manifest,'equation_count':len(EQUATIONS),
        'source_sha256':env['source_sha256'],'numbers_source':'results/summary.json'},indent=2)+'\n')
    print(json.dumps({'docx':str(target),'layout':'Continuous main text with separate code appendices','appendix_pages':len(appendix_manifest),
                       'equations':len(EQUATIONS),'largest_code_line':max(z['maximum_line_length'] for z in appendix_manifest)},indent=2))


if __name__=='__main__':
    main()
